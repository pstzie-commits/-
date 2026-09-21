"""엔지니어링 문서(PDF/DOCX/TXT)에서 텍스트를 추출한다."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path


@dataclass
class LoadedDocument:
    path: Path
    text: str
    page_count: int | None  # PDF만 페이지 수 제공, 그 외는 None


def load_document(path: str | Path) -> LoadedDocument:
    p = Path(path)
    if not p.exists():
        raise FileNotFoundError(f"문서를 찾을 수 없습니다: {p}")

    suffix = p.suffix.lower()
    if suffix == ".pdf":
        return _load_pdf(p)
    if suffix == ".docx":
        return _load_docx(p)
    if suffix in (".txt", ".md"):
        return LoadedDocument(path=p, text=p.read_text(encoding="utf-8"), page_count=None)

    raise ValueError(
        f"지원하지 않는 파일 형식입니다: {suffix} (지원 형식: .pdf, .docx, .txt, .md)"
    )


def _load_pdf(p: Path) -> LoadedDocument:
    from pypdf import PdfReader

    reader = PdfReader(str(p))
    parts = []
    for i, page in enumerate(reader.pages, start=1):
        page_text = page.extract_text() or ""
        parts.append(f"\n\n--- [PDF {i}페이지] ---\n{page_text}")
    return LoadedDocument(path=p, text="".join(parts), page_count=len(reader.pages))


def _load_docx(p: Path) -> LoadedDocument:
    import docx

    doc = docx.Document(str(p))
    parts = [para.text for para in doc.paragraphs]

    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))

    return LoadedDocument(path=p, text="\n".join(parts), page_count=None)
